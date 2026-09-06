#!/usr/bin/env python
# -*- coding: utf-8 -*-
# By: Wang Li 2023
import os, time
import argparse
import fitz
import subprocess
import pandas as pd
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import json


def table_replace(docx, value, value_re):
    for table in docx.tables:  # 遍历文档中的所有表格
        for row in table.rows:  # 遍历表格中的所有行
            for cell in row.cells:  # 遍历行中的所有单元格
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        if "%s" % value in run.text:
                            run.text = run.text.replace(value, value_re)


def conver_img(pdf):
    doc = fitz.open(pdf)
    pdf_file = os.path.split(pdf)[1]
    pdf_name = os.path.splitext(pdf_file)[0]
    page = doc[0]
    zoom_x, zoom_y = 3, 3
    mat = fitz.Matrix(zoom_x, zoom_y)
    pm = page.get_pixmap(matrix=mat)
    pm.save('%s.png' % pdf_name)
    return '%s.png' % pdf_name


def RP_img(img_ls, tag, paragraphs):
    try:
        cmd = r'ls %s' % img_ls
        img = subprocess.getoutput(cmd).split('\n')[0]
        if u'无法访问' not in img and 'cannot access' not in img:
            pic = img
            for i in range(len(paragraphs)):
                if tag in paragraphs[i].text:
                    run = paragraphs[i + 1].runs[0]
                    run.clear()
                    run.add_picture(pic, width=Cm(10))
        # elif u'无法访问' and 'upset.pdf' in img:
        # 	for i in range(len(paragraphs)):
        # 		if tag in paragraphs[i].text:
        # 			for paragraph in paragraphs[i:i + 3]:
        # 				p = paragraph._element
        # 				p.getparent().remove(p)
        else:
            print(u'\t%s\n\t\t没有这张图片' % tag)
            print(cmd)
    except Exception as e:
        print('\t%s\n\t\t RAISE ERROR' % tag)
        print(e)


def move_table_after(table, tag, docx):
    paragraphs = docx.paragraphs
    for i in range(len(paragraphs)):
        if tag in paragraphs[i].text:
            tbl, p = table._tbl, paragraphs[i]._p
            p.addnext(tbl)


def df2table(df, docx):
    if len(df.columns) == 2:
        table = docx.add_table(rows=1, cols=len(df.columns), style="Table Grid")
        table.cell(0, 0).text = "Sample ID/样本名称"
        table.cell(0, 1).text = "Group/组别"
        for _, row in df.iterrows():
            table_row = table.add_row()
            for i, val in enumerate(row):
                cell = table_row.cells[i]
                cell.text = str(val)
    elif len(df.columns) == 4:
        table = docx.add_table(rows=1, cols=len(df.columns), style="Table Grid")
        table.cell(0, 0).text = "比较组"
        table.cell(0, 1).text = "分子（实验组）"
        table.cell(0, 2).text = "分母（对照组）"
        table.cell(0, 3).text = "统计方法"
        for _, row in df.iterrows():
            table_row = table.add_row()
            for i, val in enumerate(row):
                cell = table_row.cells[i]
                cell.text = str(val)
    table.style.font.size = Pt(10)
    table.style.paragraph_format.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j in range(len(df.columns)):
        cell = table.cell(0, j)
        p = cell.paragraphs[0]
        # 设置居中
        p.paragraph_format.alignment = WD_TABLE_ALIGNMENT.CENTER
        # 设置字体大小
        p.runs[0].font.size = Pt(10)
        # 设置表头单元格填充
        bg = parse_xml(rf'<w:shd {nsdecls("w")} w:fill="2b8cbe"/>')
        cell._tc.get_or_add_tcPr().append(bg)
    return table


def get_compare(datadir):
    compare = pd.read_csv('%s/comparison.txt' % datadir, sep='\t')
    A_vs_Bs = []
    stats = []
    for index, row in compare.iterrows():
        A_vs_B = row[0] + '_vs_' + row[1]
        A_vs_Bs.append(A_vs_B)
        stats.append('T-test')
    A_vs_Bs = pd.Series(A_vs_Bs)
    stats = pd.Series(stats)
    compare.insert(0, 'vs', value=A_vs_Bs)
    compare.insert(3, 'stats', value=stats)
    return compare


def get_businfo(datadir):
    sample = pd.read_csv('%s/sample-metadata.tsv' % datadir, sep='\t', skiprows=[1])
    # bus_info = pd.read_csv('%s/bus_info.txt' % datadir, sep='\t')
    bus_info=json.load(open(os.path.join(datadir,"project_info.json"),mode='r',encoding="utf-8"))
    name = bus_info['客户名称']
    partment = bus_info['客户单位']
    NO = bus_info['项目编号']
    TI = str(time.strftime("%Y-%m-%d", time.localtime()))
    return sample, name, partment, NO, TI


def Micro_RP(datadir, resdir, docx_path, analyse, binning):
    if analyse == 'yes' and binning == 'yes':
        micro_docx = Document('%s/metagenome_megahit_bins.docx' % docx_path)
    elif analyse == 'yes' and binning == 'no':
        micro_docx = Document('%s/metagenome_megahit.docx' % docx_path)
    elif analyse == 'no':
        micro_docx = Document('%s/metagenome_megahit_Noana.docx' % docx_path)
    paragraphs = micro_docx.paragraphs
    # if os.path.exists('%s/comparison.txt' % datadir):
    # 	compare = get_compare(datadir)
    sample, name, partment, NO, TI = get_businfo(datadir)
    # s_table = df2table(sample, micro_docx)
    # c_table = df2table(compare, micro_docx)
    # move_table_after(s_table, u'本项目样本信息如下表所示', micro_docx)
    # move_table_after(c_table, u'分组方案如下表所示', micro_docx)

    table_replace(micro_docx, "姓名", name)
    table_replace(micro_docx, "检测部", partment)
    table_replace(micro_docx, "目号", NO)
    table_replace(micro_docx, "时间", TI)

    print(r'%s/group*/1-data_quality/*/error_rate.png' % resdir)
    RP_img(r'%s/group*/1-data_quality/*/error_rate.png' % resdir, u'原始数据碱基质量值分布图', paragraphs)
    RP_img(r'%s/group*/1-data_quality/*/ATGC_content.png' % resdir, u'碱基含量分布图', paragraphs)
    RP_img(r'%s/group*/1-data_quality/*/reads_quality_summary.png' % resdir, u'原始数据组成图', paragraphs)
    RP_img(r'%s/group*/2-Assembly/contig_length.png' % resdir, u'Contigs长度分布', paragraphs)
    RP_img(r'%s/group*/3-GenePredict/gene_length.png' % resdir, u'gene catalogue长度分布图', paragraphs)
    RP_img(r'%s/group*/4-GeneAbundance/Sample_correlation/sample.corr_heatmap.png' % resdir, u'样品间相关系数热图',
        paragraphs)
    RP_img(r'%s/group*/4-GeneAbundance/Venn/upset.png' % resdir, u'韦恩图(Venn Graph)或者upset图', paragraphs)
    RP_img(r'%s/group*/5-TaxAnnotation/3.Barplot/Samples/All/*.png' % resdir, u'多样本群落结构柱状图示例图', paragraphs)
    RP_img(r'%s/group*/5-TaxAnnotation/4.Bar_tree/All/*.png' % resdir, u'聚类树柱状图组合示例', paragraphs)
    RP_img(r'%s/group*/5-TaxAnnotation/5.Heatmap/Samples/All/*.png' % resdir, u'物种组成聚类热图如下', paragraphs)
    RP_img(r'%s/group*/5-TaxAnnotation/6.Beta_diversity_analysis/*/*/1.PCA/*.png' % resdir, u'PCA图如下所示',
        paragraphs)
    RP_img(r'%s/group*/5-TaxAnnotation/6.Beta_diversity_analysis/*/*/2.PCoA/*.png' % resdir, u'PCoA图如下所示',
        paragraphs)
    RP_img(r'%s/group*/5-TaxAnnotation/6.Beta_diversity_analysis/*/*/3.NMDS/*.png' % resdir, u'NMDS分析结果图如下',
        paragraphs)
    RP_img(r'%s/group*/5-TaxAnnotation/7.alpha_diversity_analysis/*/*.png' % resdir, u'样本组间alpha多样性指数箱体图如下',
        paragraphs)
    RP_img(r'%s/group*/6-TaxStatistical_analysis/*/*/1.ANOVA/*.png' % resdir, u'组间Anova方差分析柱状图', paragraphs)
    RP_img(r'%s/group*/6-TaxStatistical_analysis/*/*/2.wilcoxon/*.png' % resdir, u'组间秩和检验差异物种柱状图',
        paragraphs)
    RP_img(r'%s/group*/6-TaxStatistical_analysis/*/*/3.Stamp/*.png' % resdir, u'组间Stamp分析物种柱状图', paragraphs)
    RP_img(r'%s/group*/6-TaxStatistical_analysis/*/*/4.Random_Forest/*.png' % resdir, u'随机森林分析图如下所示',
        paragraphs)
    RP_img(r'%s/group*/6-TaxStatistical_analysis/*/*/5.metagenomeSeq/*.png' % resdir, u'MetagenomeSeq差异热图如下',
        paragraphs)
    RP_img(r'%s/group*/6-TaxStatistical_analysis/*/*/6.Anosim/*.png' % resdir, u'Anosim结果如图', paragraphs)
    RP_img(r'%s/group*/6-TaxStatistical_analysis/*/*/9.Lefse/*.png' % resdir, u'LDA值分布柱状图', paragraphs)
    RP_img(r'%s/group*/7-FunctionAnnotation/*/1.Barplot/*.png' % resdir, u'功能丰度柱形图如下', paragraphs)
    RP_img(r'%s/group*/7-FunctionAnnotation/*/2.Heatmap/*.png' % resdir, u'功能丰度热图如下', paragraphs)
    RP_img(r'%s/group*/7-FunctionAnnotation/*/3.PCA/*.png' % resdir, u'基于功能丰度的PCA分析', paragraphs)
    RP_img(r'%s/group*/7-FunctionAnnotation/*/4.PCoA/*.png' % resdir, u'基于功能丰度的PCoA分析', paragraphs)
    RP_img(r'%s/group*/7-FunctionAnnotation/*/5.NMDS/*.png' % resdir, u'基于功能丰度的NMDS分析', paragraphs)
    RP_img(r'%s/group*/8-FunctionStatistical_analysis/*/1.ANOVA/*.png' % resdir, u'基于功能丰度的ANOVA分析差异柱状图',
        paragraphs)
    RP_img(r'%s/group*/8-FunctionStatistical_analysis/*/2.wilcoxon/*.png' % resdir, u'基于功能丰度的秩和检验差异柱状图',
        paragraphs)
    RP_img(r'%s/group*/8-FunctionStatistical_analysis/*/3.Stamp/*.png' % resdir, u'基于功能丰度的stamp差异图',
        paragraphs)
    RP_img(r'%s/group*/8-FunctionStatistical_analysis/*/4.Random_Forest/*.png' % resdir, u'功能通路的随机森林分析图',
        paragraphs)
    RP_img(r'%s/group*/8-FunctionStatistical_analysis/*/5.metagenomeSeq/*.png' % resdir, u'显著差异功能热图',
        paragraphs)
    RP_img(r'%s/group*/8-FunctionStatistical_analysis/*/9.Lefse/*.png' % resdir, u'基于功能丰度的LDA柱状图', paragraphs)
    RP_img(r'%s/group*/9-METABOLIC/*/1.Barplot/*.png' % resdir, u'不同循环功能丰度柱形图', paragraphs)
    RP_img(r'%s/group*/9-METABOLIC/*/2.Heatmap/*.png' % resdir, u'不同循环功能丰度热图', paragraphs)
    RP_img(r'%s/group*/9-METABOLIC/*/3.PCA/*.png' % resdir, u'不同循环的PCA分析如下图', paragraphs)
    RP_img(r'%s/group*/9-METABOLIC/*/4.PCoA/*.png' % resdir, u'不同循环的PCoA分析如下', paragraphs)
    RP_img(r'%s/group*/9-METABOLIC/*/5.NMDS/*.png' % resdir, u'不同循环的NMDS分析如下', paragraphs)
    RP_img(r'%s/group*/9-METABOLIC/*/6.Statistical_test_analysis/1.ANOVA/*.png' % resdir,
        u'不同循环的ANOVA分析差异柱状图', paragraphs)
    RP_img(r'%s/group*/9-METABOLIC/*/6.Statistical_test_analysis/2.wilcoxon/*.png' % resdir,
        u'不同循环的秩和检验差异柱状图', paragraphs)
    RP_img(r'%s/group*/9-METABOLIC/*/6.Statistical_test_analysis/3.Stamp/*.png' % resdir, u'不同循环的stamp差异图',
        paragraphs)
    RP_img(r'%s/group*/9-METABOLIC/*/6.Statistical_test_analysis/4.Random_Forest/*.png' % resdir,
        u'不同循环的随机森林分析图', paragraphs)
    RP_img(r'%s/group*/9-METABOLIC/*/6.Statistical_test_analysis/5.metagenomeSeq/*.png' % resdir,
        u'不同循环的差异功能热图', paragraphs)
    RP_img(r'%s/group*/9-METABOLIC/*/6.Statistical_test_analysis/9.Lefse/*.png' % resdir, u'不同循环的LDA柱状图',
        paragraphs)
    RP_img(r'%s/group*/10-ARG/1.Barplot/*.png' % resdir, u'ARG丰度柱形图如下', paragraphs)
    RP_img(r'%s/group*/10-ARG/2.Heatmap/*.png' % resdir, u'ARG功能丰度聚类热图', paragraphs)
    RP_img(r'%s/group*/10-ARG/3.PCA/*.png' % resdir, u'ARG功能丰度PCA分析图', paragraphs)
    RP_img(r'%s/group*/10-ARG/4.PCoA/*.png' % resdir, u'ARG功能丰度PCoA分析图如下', paragraphs)
    RP_img(r'%s/group*/10-ARG/5.NMDS/*.png' % resdir, u'ARG功能丰度NMDS分析图', paragraphs)
    RP_img(r'%s/group*/10-ARG/6.Statistical_test_analysis/1.ANOVA/*.png' % resdir, u'ARG的ANOVA分析差异功能柱状图',
        paragraphs)
    RP_img(r'%s/group*/10-ARG/6.Statistical_test_analysis/2.wilcoxon/*.png' % resdir, u'ARG的秩和检验差异功能柱状图',
        paragraphs)
    RP_img(r'%s/group*/10-ARG/6.Statistical_test_analysis/3.Stamp/*.png' % resdir, u'ARG差异检验柱状图如下', paragraphs)
    RP_img(r'%s/group*/10-ARG/6.Statistical_test_analysis/4.Random_Forest/*.png' % resdir, u'ARG随机森林分析',
        paragraphs)
    RP_img(r'%s/group*/10-ARG/6.Statistical_test_analysis/5.metagenomeSeq/*.png' % resdir, u'ARG差异功能热图',
        paragraphs)
    RP_img(r'%s/group*/10-ARG/6.Statistical_test_analysis/9.Lefse/*.png' % resdir, u'ARG的LDA值柱状图', paragraphs)
    RP_img(r'%s/group*/11-VFDB/1.Barplot/*.png' % resdir, u'毒力因子基因丰度柱形图', paragraphs)
    RP_img(r'%s/group*/11-VFDB/2.Heatmap/*.png' % resdir, u'毒力因子基因聚类热图', paragraphs)
    RP_img(r'%s/group*/11-VFDB/3.PCA/*.png' % resdir, u'毒力因子基因PCA分析图', paragraphs)
    RP_img(r'%s/group*/11-VFDB/4.PCoA/*.png' % resdir, u'毒力因子基因PCoA分析', paragraphs)
    RP_img(r'%s/group*/11-VFDB/5.NMDS/*.png' % resdir, u'毒力因子基因NMDS分析', paragraphs)
    RP_img(r'%s/group*/11-VFDB/6.Statistical_test_analysis/1.ANOVA/*.png' % resdir, u'毒力因子基因ANOVA分析', paragraphs)
    RP_img(r'%s/group*/11-VFDB/6.Statistical_test_analysis/2.wilcoxon/*.png' % resdir, u'毒力因子基因秩和检验', paragraphs)
    RP_img(r'%s/group*/11-VFDB/6.Statistical_test_analysis/3.Stamp/*.png' % resdir, u'毒力因子基因差异检验柱状图', paragraphs)
    RP_img(r'%s/group*/11-VFDB/6.Statistical_test_analysis/4.Random_Forest/*.png' % resdir, u'毒力因子基因随机森林分析', paragraphs)
    RP_img(r'%s/group*/11-VFDB/6.Statistical_test_analysis/5.metagenomeSeq/*.png' % resdir, u'毒力因子基因差异功能热图', paragraphs)
    RP_img(r'%s/group*/11-VFDB/6.Statistical_test_analysis/9.Lefse/*.png' % resdir, u'毒力因子基因LDA值柱状图', paragraphs)

    RP_img(r'%s/group*/12-mobileOG/1.Barplot/*.png' % resdir, u'可移动基因元件丰度柱形图如下：', paragraphs)
    RP_img(r'%s/group*/12-mobileOG/2.Heatmap/*.png' % resdir, u'可移动基因元件聚类热图：', paragraphs)
    RP_img(r'%s/group*/12-mobileOG/3.PCA/*.png' % resdir, u'可移动基因元件PCA分析图：', paragraphs)
    RP_img(r'%s/group*/12-mobileOG/4.PCoA/*.png' % resdir, u'可移动基因元件PCoA分析：', paragraphs)
    RP_img(r'%s/group*/12-mobileOG/5.NMDS/*.png' % resdir, u'可移动基因元件NMDS分析：', paragraphs)
    RP_img(r'%s/group*/12-mobileOG/6.Statistical_test_analysis/1.ANOVA/*.png' % resdir, u'可移动基因元件ANOVA分析：', paragraphs)
    RP_img(r'%s/group*/12-mobileOG/6.Statistical_test_analysis/2.wilcoxon/*.png' % resdir, u'可移动基因元件秩和检验：', paragraphs)
    RP_img(r'%s/group*/12-mobileOG/6.Statistical_test_analysis/3.Stamp/*.png' % resdir, u'可移动基因元件差异检验柱状图如下所示：', paragraphs)
    RP_img(r'%s/group*/12-mobileOG/6.Statistical_test_analysis/4.Random_Forest/*.png' % resdir, u'可移动基因元件随机森林分析（排名前10）：', paragraphs)
    RP_img(r'%s/group*/12-mobileOG/6.Statistical_test_analysis/5.metagenomeSeq/*.png' % resdir, u'可移动基因元件差异功能热图如下：', paragraphs)
    RP_img(r'%s/group*/12-mobileOG/6.Statistical_test_analysis/9.Lefse/*.png' % resdir, u'可移动基因元件LDA值柱状图如下所示：', paragraphs)

    RP_img(r'%s/group*/13-BacMet2/1.Barplot/*.png' % resdir, u'重金属抗性基因丰度柱形图如下：', paragraphs)
    RP_img(r'%s/group*/13-BacMet2/2.Heatmap/*.png' % resdir, u'重金属抗性基因聚类热图：', paragraphs)
    RP_img(r'%s/group*/13-BacMet2/3.PCA/*.png' % resdir, u'重金属抗性基因PCA分析图：', paragraphs)
    RP_img(r'%s/group*/13-BacMet2/4.PCoA/*.png' % resdir, u'重金属抗性基因PCoA分析：', paragraphs)
    RP_img(r'%s/group*/13-BacMet2/5.NMDS/*.png' % resdir, u'重金属抗性基因NMDS分析：', paragraphs)
    RP_img(r'%s/group*/13-BacMet2/6.Statistical_test_analysis/1.ANOVA/*.png' % resdir, u'重金属抗性基因ANOVA分析：', paragraphs)
    RP_img(r'%s/group*/13-BacMet2/6.Statistical_test_analysis/2.wilcoxon/*.png' % resdir, u'重金属抗性基因秩和检验：', paragraphs)
    RP_img(r'%s/group*/13-BacMet2/6.Statistical_test_analysis/3.Stamp/*.png' % resdir, u'重金属抗性基因差异检验柱状图如下所示：', paragraphs)
    RP_img(r'%s/group*/13-BacMet2/6.Statistical_test_analysis/4.Random_Forest/*.png' % resdir, u'重金属抗性基因随机森林分析（排名前10）：', paragraphs)
    RP_img(r'%s/group*/13-BacMet2/6.Statistical_test_analysis/5.metagenomeSeq/*.png' % resdir, u'重金属抗性基因差异功能热图如下：', paragraphs)
    RP_img(r'%s/group*/13-BacMet2/6.Statistical_test_analysis/9.Lefse/*.png' % resdir, u'重金属抗性基因LDA值柱状图如下所示：', paragraphs)

    RP_img(r'%s/group*/14-QS/1.Barplot/*.png' % resdir, u'群体感应基因丰度柱形图如下：', paragraphs)
    RP_img(r'%s/group*/14-QS/2.Heatmap/*.png' % resdir, u'群体感应基因聚类热图：', paragraphs)
    RP_img(r'%s/group*/14-QS/3.PCA/*.png' % resdir, u'群体感应基因PCA分析图：', paragraphs)
    RP_img(r'%s/group*/14-QS/4.PCoA/*.png' % resdir, u'群体感应基因PCoA分析：', paragraphs)
    RP_img(r'%s/group*/14-QS/5.NMDS/*.png' % resdir, u'群体感应基因NMDS分析：', paragraphs)
    RP_img(r'%s/group*/14-QS/6.Statistical_test_analysis/1.ANOVA/*.png' % resdir, u'群体感应基因ANOVA分析：', paragraphs)
    RP_img(r'%s/group*/14-QS/6.Statistical_test_analysis/2.wilcoxon/*.png' % resdir, u'群体感应基因秩和检验：', paragraphs)
    RP_img(r'%s/group*/14-QS/6.Statistical_test_analysis/3.Stamp/*.png' % resdir, u'群体感应基因差异检验柱状图如下所示：', paragraphs)
    RP_img(r'%s/group*/14-QS/6.Statistical_test_analysis/4.Random_Forest/*.png' % resdir, u'群体感应基因随机森林分析（排名前10）：', paragraphs)
    RP_img(r'%s/group*/14-QS/6.Statistical_test_analysis/5.metagenomeSeq/*.png' % resdir, u'群体感应基因差异功能热图如下：', paragraphs)
    RP_img(r'%s/group*/14-QS/6.Statistical_test_analysis/9.Lefse/*.png' % resdir, u'群体感应基因LDA值柱状图如下所示：', paragraphs)

    RP_img(r'%s/binning/2.Bin_Plot/*.png' % resdir, u'分箱GC_coverage图：', paragraphs)
    RP_img(r'%s/binning/3.Bin_Abundance/*.png' % resdir, u'分箱丰度聚类热图：', paragraphs)

    micro_docx.save('%s/report.docx' % resdir)


def main():
    parser = argparse.ArgumentParser(
        description='This script will generate Result report')
    parser.add_argument('-I', '--i_datadir', type=str, required=True, help='the dir of sample-metadata.tsv')
    parser.add_argument('--analyse', type=str, choices=['yes', 'no'], help='Whether to analyze')
    parser.add_argument('--binning', type=str, choices=['yes', 'no'], help='Whether to binning')
    parser.add_argument('--res_dir', type=str, required=True, help='the dir of Result')
    parser.add_argument('--micro_docx_path', type=str, default='/home/wangli/microbiome/microbiome',
                        help='the dir of micro_docx_path')
    args = parser.parse_args()
    datadir = os.path.abspath(args.i_datadir)
    res_dir = os.path.abspath(args.res_dir)
    docx_path = os.path.abspath(args.micro_docx_path)
    analyse = args.analyse
    binning = args.binning

    Micro_RP(datadir, res_dir, docx_path, analyse, binning)
    os.system("libreoffice7.5 --headless --convert-to pdf --outdir %s %s/report.docx" % (res_dir, res_dir))


if __name__ == '__main__':
    main()
